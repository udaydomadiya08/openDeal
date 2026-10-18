"""PEAK input mapping registry and authorization whitelist."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import openpyxl

from opendeal.valuation.contracts import ValuationCellDefinition
from opendeal.valuation.exceptions import AmbiguousMappingError, FormulaOverwriteProhibitedError


class PEAKInputRegistry:
    """Authoritative whitelist of cells explicitly authorized as valuation inputs.

    Rejects any cell write to formulas or unmapped coordinates.
    """

    def __init__(self, mappings: List[ValuationCellDefinition], version: str = "2.0.0") -> None:
        self.version = version
        self._by_coordinate: Dict[str, ValuationCellDefinition] = {}
        self._by_variable: Dict[str, ValuationCellDefinition] = {}

        for m in mappings:
            # Formula cells are strictly prohibited as input cells
            if m.is_formula:
                raise FormulaOverwriteProhibitedError(
                    f"Cell {m.coordinate_key} contains a formula and cannot be registered as an input."
                )
            coord = m.coordinate_key.upper()
            if coord in self._by_coordinate:
                # Disallow ambiguous conflicting mappings
                existing = self._by_coordinate[coord]
                if existing.canonical_variable != m.canonical_variable:
                    raise AmbiguousMappingError(
                        f"Ambiguous mapping collision for {coord}: '{existing.canonical_variable}' vs '{m.canonical_variable}'"
                    )
            self._by_coordinate[coord] = m
            self._by_variable[m.canonical_variable] = m

    @property
    def total_authorized_cells(self) -> int:
        return len(self._by_coordinate)

    def is_authorized(self, sheet: str, cell: str) -> bool:
        """Check if sheet!cell is authorized for writing."""
        coord = f"{sheet}!{cell}".upper()
        return coord in self._by_coordinate

    def get_definition(self, sheet: str, cell: str) -> Optional[ValuationCellDefinition]:
        """Get cell definition by coordinate."""
        coord = f"{sheet}!{cell}".upper()
        return self._by_coordinate.get(coord)

    def get_by_variable(self, canonical_variable: str) -> Optional[ValuationCellDefinition]:
        """Get cell definition by canonical variable name."""
        return self._by_variable.get(canonical_variable)

    def all_definitions(self) -> List[ValuationCellDefinition]:
        """Return all authorized cell definitions."""
        return list(self._by_coordinate.values())

    def verify_against_workbook(self, workbook_path: Path) -> Tuple[bool, List[str]]:
        """Verify that every registered input cell actually exists and is NOT a formula."""
        wb = openpyxl.load_workbook(workbook_path, data_only=False, read_only=True)
        errors = []
        for defn in self._by_coordinate.values():
            if defn.sheet not in wb.sheetnames:
                errors.append(f"Sheet '{defn.sheet}' does not exist in workbook.")
                continue
            ws = wb[defn.sheet]
            c = ws[defn.cell]
            val = c.value
            if isinstance(val, str) and val.startswith("="):
                errors.append(
                    f"Cell {defn.coordinate_key} contains formula '{val}' but was registered as an input."
                )
        wb.close()
        return (len(errors) == 0, errors)

    @classmethod
    def load_from_json(cls, json_path: Path, version: str = "2.0.0") -> PEAKInputRegistry:
        """Load input mappings from JSON file."""
        with open(json_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        definitions = []
        for item in data:
            defn = ValuationCellDefinition(
                canonical_variable=item.get("var", ""),
                sheet=item.get("sheet", ""),
                cell=item.get("cell", ""),
                unit=item.get("unit", ""),
                description=item.get("desc", ""),
                is_formula=False,
            )
            definitions.append(defn)

        return cls(definitions, version=version)

    @classmethod
    def default(cls) -> PEAKInputRegistry:
        """Load default verified PEAK input mappings bundled with OpenDeal."""
        json_path = Path(__file__).parent / "verified_mappings.json"
        if not json_path.exists():
            raise FileNotFoundError(f"Verified mappings file not found: {json_path}")
        return cls.load_from_json(json_path, version="2.0.0")
