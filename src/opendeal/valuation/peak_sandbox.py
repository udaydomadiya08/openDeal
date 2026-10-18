"""PEAK Sandbox Runner — Isolated, verified execution of PEAK workbook valuations."""

from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
from typing import Any, Dict, List, Optional, Tuple, Union
import uuid
import openpyxl

from opendeal.valuation.contracts import (
    CalculationMode,
    CalculationStatus,
    CellWriteAudit,
    RunManifest,
    ValuationInputPayload,
    WorkbookIntegrityReport,
)
from opendeal.valuation.exceptions import (
    AmbiguousMappingError,
    FormulaOverwriteProhibitedError,
    UnauthorizedCellWriteError,
    WorkbookIntegrityError,
)
from opendeal.valuation.mapping import PEAKInputRegistry


def compute_file_sha256(path: Path) -> str:
    """Compute the SHA-256 checksum of a file in binary mode."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


class PEAKSandboxRunner:
    """Manages isolated, copy-on-write valuation runs on the institutional PEAK workbook.

    Enforces strict immutability of the golden master workbook, prevents writing to
    unauthorized cells or formula cells, and produces comprehensive audit manifests.
    """

    def __init__(
        self,
        master_workbook_path: Union[str, Path],
        runs_root_dir: Union[str, Path],
        registry: Optional[PEAKInputRegistry] = None,
    ) -> None:
        self.master_workbook_path = Path(master_workbook_path).resolve()
        if not self.master_workbook_path.exists():
            raise FileNotFoundError(f"Golden master workbook not found at: {self.master_workbook_path}")

        self.runs_root_dir = Path(runs_root_dir).resolve()
        self.runs_root_dir.mkdir(parents=True, exist_ok=True)

        self.registry = registry or PEAKInputRegistry.default()

    def snapshot_formulas(self, wb: openpyxl.Workbook) -> Dict[str, str]:
        """Extract all formula expressions from the workbook indexed by Sheet!Cell."""
        formulas: Dict[str, str] = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(values_only=False):
                for cell in row:
                    val = cell.value
                    if isinstance(val, str) and val.startswith("="):
                        formulas[f"{sheet_name}!{cell.coordinate}"] = val
        return formulas

    def execute_run(
        self,
        inputs: List[ValuationInputPayload],
        entity_id: str,
        run_id: Optional[str] = None,
        mode: CalculationMode = CalculationMode.HEADLESS,
    ) -> RunManifest:
        """Execute an isolated sandbox valuation run.

        1. Verifies master workbook integrity before starting.
        2. Copies master workbook into an isolated run directory.
        3. Enforces that no writes target formula cells or unauthorized coordinates.
        4. Applies validated inputs.
        5. Verifies formula preservation and structure post-write.
        6. Confirms the golden master workbook remains bit-for-bit unchanged.
        7. Produces a machine-readable run manifest.
        """
        run_start = datetime.now(timezone.utc)
        if not run_id:
            run_id = f"run_{run_start.strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}"
        else:
            if not re.match(r"^[a-zA-Z0-9_\-]+$", run_id):
                raise ValueError(f"Invalid run_id '{run_id}': Directory traversal prohibited.")

        candidate_run_dir = self.runs_root_dir / run_id
        if candidate_run_dir.is_symlink():
            raise ValueError(f"Invalid run_id '{run_id}': Target is a symbolic link.")

        resolved_runs_root = self.runs_root_dir.resolve()
        run_dir = candidate_run_dir.resolve()
        if not run_dir.is_relative_to(resolved_runs_root) or run_dir == resolved_runs_root:
            raise ValueError(f"Invalid run_id '{run_id}': Directory traversal prohibited.")

        run_dir.mkdir(parents=True, exist_ok=True)

        sandbox_wb_path = run_dir / "valuation_sandbox.xlsx"
        manifest_path = run_dir / "run_manifest.json"

        # Prevent following symlinks on write
        if sandbox_wb_path.is_symlink():
            raise WorkbookIntegrityError(f"Target workbook path '{sandbox_wb_path}' is a symbolic link; write aborted.")
        if manifest_path.is_symlink():
            raise WorkbookIntegrityError(f"Target manifest path '{manifest_path}' is a symbolic link; write aborted.")

        # 1. Hash golden master before copying
        master_sha256_pre = compute_file_sha256(self.master_workbook_path)

        # 2. Copy master to isolated run directory
        shutil.copyfile(self.master_workbook_path, sandbox_wb_path, follow_symlinks=False)
        sandbox_sha256_pre = compute_file_sha256(sandbox_wb_path)

        if sandbox_sha256_pre != master_sha256_pre:
            raise WorkbookIntegrityError("Sandboxed copy hash does not match master before execution.")

        # 3. Load sandbox workbook
        wb = openpyxl.load_workbook(sandbox_wb_path, data_only=False)

        # 4. Snapshot existing formulas
        formulas_pre = self.snapshot_formulas(wb)
        formula_count_pre = len(formulas_pre)

        # 5. Pre-validate all write targets
        for inp in inputs:
            coord = inp.coordinate_key.upper()
            sheet = inp.target_sheet
            cell = inp.target_cell

            # Check sheet exists
            if sheet not in wb.sheetnames:
                raise AmbiguousMappingError(f"Target sheet '{sheet}' does not exist in workbook.")

            # Check authorization whitelist
            if not self.registry.is_authorized(sheet, cell):
                raise UnauthorizedCellWriteError(
                    f"Cell '{coord}' is not authorized as a PEAK input cell in the whitelist registry."
                )

            # Check target cell is not currently a formula
            current_val = wb[sheet][cell].value
            if isinstance(current_val, str) and current_val.startswith("="):
                raise FormulaOverwriteProhibitedError(
                    f"Cell '{coord}' contains formula '{current_val}'. Writing to formula cells is prohibited."
                )

        # 6. Apply authorized inputs
        applied_writes: List[CellWriteAudit] = []
        for inp in inputs:
            ws = wb[inp.target_sheet]
            c = ws[inp.target_cell]
            prev_val = c.value
            c.value = inp.value
            applied_writes.append(
                CellWriteAudit(
                    target_sheet=inp.target_sheet,
                    target_cell=inp.target_cell,
                    canonical_variable=inp.canonical_variable,
                    previous_value=prev_val,
                    new_value=inp.value,
                    fact_id=inp.fact_id,
                    applied_at=datetime.now(timezone.utc),
                    status="APPLIED",
                )
            )

        # 7. Save sandbox workbook
        wb.save(sandbox_wb_path)
        wb.close()

        sandbox_sha256_post = compute_file_sha256(sandbox_wb_path)

        # 8. Re-open and verify formula preservation
        wb_verify = openpyxl.load_workbook(sandbox_wb_path, data_only=False)
        formulas_post = self.snapshot_formulas(wb_verify)
        formula_count_post = len(formulas_post)

        formula_mismatches: List[str] = []
        for coord, expr in formulas_pre.items():
            if coord not in formulas_post:
                formula_mismatches.append(f"Missing formula at {coord}: expected '{expr}'")
            elif formulas_post[coord] != expr:
                formula_mismatches.append(
                    f"Formula altered at {coord}: was '{expr}', now '{formulas_post[coord]}'"
                )

        wb_verify.close()

        # 9. Verify golden master is still bit-for-bit identical
        master_sha256_post = compute_file_sha256(self.master_workbook_path)
        master_unmodified = (master_sha256_pre == master_sha256_post)
        if not master_unmodified:
            raise WorkbookIntegrityError(
                "CRITICAL: Golden master workbook was modified during sandbox execution!"
            )

        formulas_identical = (len(formula_mismatches) == 0 and formula_count_pre == formula_count_post)
        if not formulas_identical:
            raise WorkbookIntegrityError(
                f"Formula integrity violation: {len(formula_mismatches)} mismatches detected."
            )

        integrity_report = WorkbookIntegrityReport(
            original_workbook_sha256_pre=master_sha256_pre,
            original_workbook_sha256_post=master_sha256_post,
            original_workbook_unmodified=master_unmodified,
            sandboxed_workbook_sha256_pre=sandbox_sha256_pre,
            sandboxed_workbook_sha256_post=sandbox_sha256_post,
            formula_count_pre=formula_count_pre,
            formula_count_post=formula_count_post,
            formulas_identical=formulas_identical,
            formula_mismatches=formula_mismatches,
            unauthorized_writes_detected=False,
            authorized_writes_count=len(applied_writes),
        )

        # 10. Handle Calculation Mode
        calc_status, calc_msg, extracted_outputs = self._handle_calculation(
            sandbox_wb_path=sandbox_wb_path, mode=mode
        )

        manifest = RunManifest(
            run_id=run_id,
            created_at=run_start,
            completed_at=datetime.now(timezone.utc),
            entity_id=entity_id,
            original_workbook_path=str(self.master_workbook_path),
            sandboxed_workbook_path=str(sandbox_wb_path),
            mapping_version=self.registry.version,
            calculation_mode=mode,
            calculation_status=calc_status,
            calculation_message=calc_msg,
            inputs_applied=applied_writes,
            integrity_report=integrity_report,
            extracted_outputs=extracted_outputs,
            errors=[],
        )

        # Write manifest to disk
        if manifest_path.is_symlink():
            raise WorkbookIntegrityError(f"Target manifest path '{manifest_path}' is a symbolic link; write aborted.")
        flags = os.O_CREAT | os.O_WRONLY | os.O_TRUNC | getattr(os, "O_NOFOLLOW", 0)
        fd = os.open(str(manifest_path), flags, 0o644)
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write(manifest.model_dump_json(indent=2))

        return manifest

    def _handle_calculation(
        self, sandbox_wb_path: Path, mode: CalculationMode
    ) -> Tuple[CalculationStatus, str, Dict[str, Any]]:
        """Handles formula recalculation according to specified mode."""
        if mode == CalculationMode.HEADLESS:
            msg = (
                "Headless mode: Inputs safely populated and formulas verified intact. "
                "Excel formulas cannot be calculated by openpyxl without an external spreadsheet calculation engine. "
                "Recalculation requires Microsoft Excel automation or a validated deterministic valuation model. "
                "Cached formula values have NOT been extracted to prevent silent stale reporting."
            )
            return (CalculationStatus.NOT_CALCULATED_HEADLESS, msg, {})

        elif mode == CalculationMode.EXCEL_DESKTOP:
            try:
                import xlwings as xw  # type: ignore

                app = xw.App(visible=False, add_book=False)
                try:
                    wb = app.books.open(str(sandbox_wb_path))
                    app.calculate()
                    wb.save()

                    # Extract representative output metrics if present
                    outputs: Dict[str, Any] = {}
                    if "Global_Executive_Dashboard" in [s.name for s in wb.sheets]:
                        sheet = wb.sheets["Global_Executive_Dashboard"]
                        outputs["dcf_equity_value"] = sheet.range("B7").value
                        outputs["valuation_summary"] = sheet.range("B10").value

                    wb.close()
                    msg = "Workbook recalculated successfully via Microsoft Excel automation."
                    return (CalculationStatus.CALCULATED_DESKTOP, msg, outputs)
                finally:
                    app.quit()
            except Exception as e:
                msg = f"Excel desktop recalculation failed or was unavailable: {str(e)}"
                return (CalculationStatus.CALCULATION_FAILED, msg, {})

        return (CalculationStatus.CALCULATION_FAILED, "Unknown calculation mode requested.", {})
